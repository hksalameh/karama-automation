from decimal import Decimal
import pandas as pd

from kafala_compare_app_v2 import (
    build_review_candidates,
    detail_total_including_duplicates,
    find_printed_total_in_raw,
)


def row_site(nat, name, amount='24'):
    return {'nat_id': nat, 'name_site': name, 'amount_site': Decimal(amount)}


def row_ref(nat, name, amount='24'):
    return {'nat_id': nat, 'name_ref': name, 'amount_ref': Decimal(amount)}


# عبدالله / عبد الله مع رقم مختلف: يجب حمايته من التصفير.
site = pd.DataFrame([row_site('2001234567', 'محمد عبدالله احمد')])
ref = pd.DataFrame([row_ref('2001234568', 'محمد عبد الله احمد')])
review, s_ids, r_ids = build_review_candidates(site, ref)
assert not review.empty
assert '2001234567' in s_ids and '2001234568' in r_ids

# خطأ حرف بالاسم + خانة واحدة بالرقم الوطني: مراجعة فقط.
site = pd.DataFrame([row_site('2012345678', 'احمد محمد الزعبي')])
ref = pd.DataFrame([row_ref('2012345679', 'احمد محمذ الزعبي')])
review, s_ids, r_ids = build_review_candidates(site, ref)
assert not review.empty
assert '2012345678' in s_ids and '2012345679' in r_ids

# شخص مختلف فعلاً: لا يجب اختراع مطابقة.
site = pd.DataFrame([row_site('3000000001', 'سليم خالد العلي')])
ref = pd.DataFrame([row_ref('4000000002', 'نور محمود الخطيب')])
review, s_ids, r_ids = build_review_candidates(site, ref)
assert review.empty and not s_ids and not r_ids

# اسم مكرر في الرعاية: يجب إظهار أكثر من احتمال وعدم الحسم آلياً.
site = pd.DataFrame([row_site('5000000001', 'محمد احمد علي')])
ref = pd.DataFrame([
    row_ref('5000000011', 'محمد احمد علي'),
    row_ref('5000000022', 'محمد احمد علي'),
])
review, s_ids, r_ids = build_review_candidates(site, ref)
assert len(review) == 2
assert review['الحالة'].str.contains('مكرر').all()

# يجب تجاهل مجموع الصفحة واختيار المجموع الإجمالي المطبوع.
raw = pd.DataFrame([
    ['', 'مجموع الصفحة', '412', ''],
    ['', '', '', ''],
    ['', 'المجموع الإجمالي', '7633.3', ''],
])
assert find_printed_total_in_raw(raw) == Decimal('7633.3')

# مجموع التفاصيل يجب أن يشمل الصفوف المكررة كلها، لا أول سجل فقط.
ref_unique = pd.DataFrame([
    row_ref('1', 'أ', '10'),
    row_ref('2', 'ب', '20'),
])
ref_dup = pd.DataFrame([
    row_ref('2', 'ب', '20'),
    row_ref('2', 'ب', '4'),
])
assert detail_total_including_duplicates(ref_unique, ref_dup) == Decimal('34')

print('matching and total safety tests passed')

# Review data must stay accessible in-app without changing the workbook or IDs.
import os
import tempfile
from pathlib import Path
from openpyxl import Workbook
from auto_update_gui import (read_review_cases, ReviewCasesWindow,
                             save_review_decisions, confirm_review_cases)

with tempfile.TemporaryDirectory() as folder:
    path = Path(folder) / 'review.xlsx'
    wb = Workbook()
    ws = wb.active
    ws.title = 'مراجعة مطابقة محتملة'
    ws.append(['رقم_وطني_كرامة', 'الاسم_في_كرامة', 'رقم_وطني_الرعاية', 'الحالة',
               'مبلغ_كرامة', 'مبلغ_الرعاية', 'الاسم_في_الرعاية'])
    ws.append(['0012345678', 'محمد عبدالله احمد', '0012345679', 'اختلاف الرقم الوطني',
               '20', '30', 'محمد عبد الله احمد'])
    ws.append(['0098765432', 'نور خالد احمد', '0098765433', 'تشابه الاسم',
               '40', '50', 'نور خالد احمد'])
    ws.append([None, None, None, None])
    dup = wb.create_sheet('تكرار ببرنامج الرعاية')
    dup.append(['nat_id', 'name_ref', 'amount_ref'])
    dup.append(['0012345679', 'محمد عبدالله احمد', 24])
    summary = wb.create_sheet('ملخص')
    summary.append(['البند', 'القيمة'])
    summary.append(['سلامة_إجمالي_الرعاية', 'لا - يوجد فرق'])
    summary.append(['المجموع_المتوقع_بعد_التعديل_الآلي', '60'])
    summary.append(['حالات_مراجعة_مطابقة_محتملة', 2])
    updates = wb.create_sheet('يحتاج تعديل')
    updates.append(['الرقم_الوطني', 'الاسم_في_الموقع', 'الاسم_في_برنامج_الرعاية',
                    'المبلغ_في_الموقع', 'المبلغ_الفعلي', 'سبب', 'تعليمات'])
    wb.save(path)
    original = path.read_bytes()
    cases = read_review_cases(path)
    assert len(cases['مراجعة مطابقة محتملة']) == 2
    assert cases['مراجعة مطابقة محتملة'][0]['رقم_وطني_كرامة'] == '0012345678'
    assert cases['تكرار ببرنامج الرعاية'][0]['nat_id'] == '0012345679'
    assert cases['ملخص'][0]['القيمة'] == 'لا - يوجد فرق'
    assert path.read_bytes() == original

    # Approval enters only the selected case; deferral stays protected and audited.
    from unittest.mock import patch
    with patch('auto_update_gui.messagebox.askyesnocancel', side_effect=[True, False]) as ask:
        reviewed_path = confirm_review_cases(None, str(path), folder)
        assert ask.call_count == 2
    reviewed_cases = read_review_cases(reviewed_path)
    assert len(reviewed_cases['مراجعة مطابقة محتملة']) == 2
    assert len(reviewed_cases['يحتاج تعديل']) == 1
    assert reviewed_cases['يحتاج تعديل'][0]['الرقم_الوطني'] == '0012345678'
    assert reviewed_cases['يحتاج تعديل'][0]['المبلغ_الفعلي'] == '30'
    audit = reviewed_cases['سجل المراجعة']
    assert [row['القرار'] for row in audit] == ['معتمد للإدخال', 'مؤجل - لم يعتمد للإدخال']
    values = {row['البند']: row['القيمة'] for row in reviewed_cases['ملخص']}
    assert Decimal(values['المجموع_المتوقع_بعد_التعديل_الآلي']) == Decimal('70')
    assert values['حالات_مراجعة_مطابقة_محتملة'] == '1'
    assert path.read_bytes() == original
    with patch('auto_update_gui.messagebox.askyesnocancel', return_value=None) as ask:
        assert confirm_review_cases(None, reviewed_path, folder) is None
        assert ask.call_count == 1  # approved case is not added again
    try:
        save_review_decisions(reviewed_path, [(cases['مراجعة مطابقة محتملة'][0], True)], folder)
        raise AssertionError('duplicate approval was accepted')
    except ValueError:
        pass

    # The Windows build runner can exercise real Tk widgets and clipboard.
    if os.name == 'nt' or os.environ.get('DISPLAY'):
        import tkinter as tk
        from kafala_compare_app_v2 import EnhancedApp
        from auto_update_gui import AutoUpdateGUI
        root = tk.Tk()
        root.withdraw()
        try:
            app = EnhancedApp(root)
            window = ReviewCasesWindow(root, cases)
            root.update_idletasks()
            assert len(window.tree.get_children()) == 2
            window.move(1)
            assert window.tree.selection() == ('1',)
            window.copy_number('site')
            assert root.clipboard_get() == '0098765432'
            window.copy_number('care')
            assert root.clipboard_get() == '0098765433'
            window.query.set('محمد')
            assert len(window.tree.get_children()) == 1
            assert '0012345678' in window.details.get('1.0', tk.END)
            window.query.set('لا يوجد هذا الاسم')
            assert len(window.tree.get_children()) == 0
            window.move(1)
            window.query.set('')
            window.category.set('تكرار ببرنامج الرعاية')
            window.refresh()
            window.copy_number('care')
            assert root.clipboard_get() == '0012345679'
            panel = tk.Frame(root)
            automation = AutoUpdateGUI(panel, str(path), {})
            # Reset clears only session data and preserves a completed browser.
            from unittest.mock import Mock
            app.ref_var.set('old-reference.xlsx')
            app.site_var.set('old-site.xlsx')
            app.out_var.set(str(path))
            app.needs_update_df = pd.DataFrame([{'الرقم_الوطني': '0012345678'}])
            app.set_review_mode('all')
            app.update_app = automation
            app.update_frame = panel
            automation.options['on_new_comparison'] = app.reset_comparison
            automation.running = True
            with patch('auto_update_gui.messagebox.showwarning') as warning:
                automation.new_comparison()
                assert warning.called
                assert app.ref_var.get() == 'old-reference.xlsx'
            automation.completed_temp_save = True
            automation.process = Mock()
            automation.process.poll.return_value = None
            automation.new_comparison()
            assert not app.ref_var.get() and not app.site_var.get() and not app.out_var.get()
            assert app.review_df.empty and not app.tree.get_children()
            assert not app.v_nat.get() and not app.v_reason.get()
            assert panel.winfo_exists()
            automation.process.terminate.assert_not_called()
            assert not hasattr(app, 'update_app')
            window.win.destroy()
        finally:
            root.destroy()
    assert path.read_bytes() == original
print('in-app review data tests passed')
